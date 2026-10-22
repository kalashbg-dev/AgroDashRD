import io
import json
from collections import Counter
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from agrodash.catalog.models import BaseUnit, Market, MarketKind, Product
from agrodash.ingestion.models import (
    IngestionRun,
    IngestionStatus,
    IngestionTrigger,
    RejectionReason,
    SourceLayout,
)
from agrodash.ingestion.service import MAX_FILE_BYTES, IngestionError, ingest_file
from agrodash.prices.models import Price, PriceSource, PriceStatus, PriceType

FIXTURES = Path(__file__).parent / "fixtures"


def load(session, name: str, content: bytes | None = None) -> IngestionRun:
    content = (FIXTURES / name).read_bytes() if content is None else content
    return ingest_file(session, content, name, IngestionTrigger.FILE)


def reasons(run: IngestionRun) -> Counter:
    return Counter(row.reason for row in run.quarantined_rows)


def test_wholesale_csv_loads_prices_and_quarantines_bad_rows(session):
    run = load(session, "mayorista.csv")

    assert (run.status, run.layout) == (IngestionStatus.SUCCEEDED, SourceLayout.WHOLESALE)
    assert (run.rows_read, run.prices_loaded, run.rows_rejected, run.values_missing) == (
        12,
        5,
        6,
        1,
    )
    assert reasons(run) == {
        RejectionReason.UNIT_MISMATCH: 1,
        RejectionReason.UNKNOWN_UNIT: 1,
        RejectionReason.INVALID_DATE: 1,
        RejectionReason.INVALID_PRICE: 1,
        RejectionReason.DUPLICATE_PRICE: 1,
        RejectionReason.UNEXPECTED_LAYOUT: 1,
    }
    rice = session.query(Product).filter_by(name="Arroz Selecto/ primera").one()
    first_week = session.query(Price).filter_by(product=rice, observed_on=date(2017, 8, 1)).one()
    assert (first_week.amount, first_week.price_per_base_unit) == (
        Decimal("2350.00"),
        Decimal("18.8000"),
    )
    assert first_week.market.name == "Mercado Nuevo"
    assert (first_week.price_type, first_week.source, first_week.status) == (
        PriceType.WHOLESALE,
        PriceSource.OFFICIAL,
        PriceStatus.VALIDATED,
    )
    assert first_week.ingestion_run_id == run.id
    assert session.query(Price).filter_by(product=rice).count() == 2
    # The type column says "Mayo+B32868:C32949rista", but the row itself is sound.
    assert session.query(Product).filter_by(name="Leche (Líquida)").one()


def test_quarantined_rows_keep_their_cells_and_line_number(session):
    run = load(session, "mayorista.csv")

    unknown = next(row for row in run.quarantined_rows if row.reason == "unknown_unit")
    assert unknown.row_number == 4
    assert json.loads(unknown.raw_values)[:4] == [
        "Ajo",
        "Legumbres-Hortalizas",
        "Mayorista",
        "Saco",
    ]
    assert "Saco" in unknown.detail


def test_retail_sheet_loads_one_price_per_market(session):
    run = load(session, "minorista.csv")

    assert (run.layout, run.prices_loaded, run.rows_rejected, run.values_missing) == (
        SourceLayout.RETAIL,
        13,
        1,
        14,
    )
    assert reasons(run) == {RejectionReason.INVALID_PRICE: 1}
    markets = {market.name: market.kind for market in session.query(Market)}
    assert len(markets) == 7
    assert markets["Supermercado"] is MarketKind.SUPERMARKET
    assert markets["Mercado Conaprope"] is MarketKind.RETAIL
    supermarket = session.query(Price).join(Market).filter(Market.name == "Supermercado").all()
    assert {(price.observed_on, price.amount) for price in supermarket} == {
        (date(2017, 8, 1), Decimal("24.50")),
        (date(2026, 9, 29), Decimal("76.48")),
    }


def test_wholesale_and_retail_names_share_the_product_catalog(session):
    load(session, "mayorista.csv")
    load(session, "minorista.csv")

    rice = session.query(Product).filter_by(name="Arroz Selecto/ primera").one()
    types = Counter(price.price_type for price in session.query(Price).filter_by(product=rice))
    assert types == {PriceType.WHOLESALE: 2, PriceType.RETAIL: 4}


def test_each_price_type_keeps_its_own_base_unit(session):
    load(session, "mayorista.csv")
    load(session, "minorista.csv")

    plantain = session.query(Product).filter_by(name="Plátano (Barahonero), primera").one()
    prices = session.query(Price).filter_by(product=plantain)
    bases = {price.price_type: price.unit.base_unit for price in prices}
    assert bases == {PriceType.WHOLESALE: BaseUnit.UNIT, PriceType.RETAIL: BaseUnit.POUND}


def test_loading_the_same_file_twice_changes_nothing(session):
    first = load(session, "mayorista.csv")
    second = load(session, "mayorista.csv")

    assert second.id == first.id
    assert session.query(IngestionRun).count() == 1
    assert session.query(Price).count() == 5


def test_the_same_data_in_another_format_only_adds_new_prices(session):
    load(session, "minorista.csv")

    run = load(session, "minorista.xls")

    assert (run.status, run.prices_loaded, run.rows_rejected) == (IngestionStatus.SUCCEEDED, 0, 1)
    assert session.query(Price).count() == 13


def test_xlsx_workbooks_are_read(session):
    workbook = openpyxl.Workbook()
    workbook.active.append(
        [
            "Arroz Selecto/ primera",
            "Cereal",
            "Mayorista",
            "Saco/125 lb",
            2350,
            "Primera",
            "Agosto",
            2017,
        ]
    )
    content = io.BytesIO()
    workbook.save(content)

    run = load(session, "mayorista.xlsx", content.getvalue())

    assert (run.layout, run.prices_loaded) == (SourceLayout.WHOLESALE, 1)


@pytest.mark.parametrize(
    ("name", "content", "message"),
    [
        ("precios.pdf", b"%PDF-1.4", "Unsupported file type '.pdf'"),
        ("precios.xlsx", b"not a workbook", "Could not read precios.xlsx"),
        ("vacio.csv", b"", "precios.csv has no rows"),
    ],
)
def test_unreadable_files_fail_with_a_recorded_run(session, name, content, message):
    message = message.replace("precios.csv", name)

    with pytest.raises(IngestionError, match=message):
        load(session, name, content)

    run = session.query(IngestionRun).one()
    assert (run.status, run.file_name, run.prices_loaded) == (IngestionStatus.FAILED, name, 0)
    assert message in run.error
    assert run.finished_at is not None
    assert session.query(Price).count() == 0


def test_a_failed_file_can_be_loaded_again_after_fixing_the_cause(session, monkeypatch):
    content = (FIXTURES / "mayorista.csv").read_bytes()
    monkeypatch.setattr("agrodash.ingestion.service.MAX_FILE_BYTES", 10)
    with pytest.raises(IngestionError, match="larger than"):
        load(session, "mayorista.csv", content)
    monkeypatch.setattr("agrodash.ingestion.service.MAX_FILE_BYTES", MAX_FILE_BYTES)

    run = load(session, "mayorista.csv", content)

    assert run.status is IngestionStatus.SUCCEEDED
    assert [r.status for r in session.query(IngestionRun).order_by(IngestionRun.id)] == [
        IngestionStatus.FAILED,
        IngestionStatus.SUCCEEDED,
    ]


RETAIL_HEADER = (
    "Productos;Rubro;Categoria;Unidad;Mercado Nuevo;Mercado Conaprope;Mercado Los Mina;"
    "Mercado Villa Consuelo;Mercado Cristo Rey;Mercadom;Supermercado;Semana;Mes;Año\r\n"
)


def test_rows_rejected_counts_rows_not_cells(session):
    content = (
        RETAIL_HEADER + "Arroz;Cereal;Minorista;lb;-1;-2;30;N/D;N/D;N/D;N/D;Primera;Agosto;2017\r\n"
    )

    run = load(session, "minorista.csv", content.encode("latin-1"))

    assert (run.prices_loaded, run.rows_rejected, len(run.quarantined_rows)) == (1, 1, 2)


def test_a_short_retail_header_fails_with_a_clear_message(session):
    content = "Productos;Rubro;Categoria;Unidad;Mercado Nuevo\r\nArroz;Cereal;Minorista;lb;20\r\n"

    with pytest.raises(IngestionError, match="retail header needs 7 market columns"):
        load(session, "minorista.csv", content.encode())


def test_a_file_where_no_row_loads_fails_and_can_be_retried(session):
    content = b"Arroz;Cereal;Mayorista;Saco;100;Primera;Agosto;2017\r\n"

    with pytest.raises(IngestionError, match="No row could be loaded"):
        load(session, "mayorista.csv", content)

    run = session.query(IngestionRun).one()
    assert run.status is IngestionStatus.FAILED
    assert session.query(Product).count() == 0


def test_a_failure_mid_load_stores_nothing_but_the_failed_run(session, monkeypatch):
    def broken_market(*_):
        raise RuntimeError("disk full")

    monkeypatch.setattr("agrodash.catalog.resolver.CatalogResolver.market", broken_market)

    with pytest.raises(IngestionError, match="Unexpected error: RuntimeError: disk full"):
        load(session, "mayorista.csv")

    assert session.query(IngestionRun).one().status is IngestionStatus.FAILED
    assert (session.query(Price).count(), session.query(Product).count()) == (0, 0)


def test_workbooks_that_unpack_too_large_are_refused(session, monkeypatch):
    workbook = openpyxl.Workbook()
    workbook.active.append(["x" * 1000])
    content = io.BytesIO()
    workbook.save(content)
    monkeypatch.setattr("agrodash.ingestion.readers.MAX_UNPACKED_BYTES", 1000)

    with pytest.raises(IngestionError, match="too large when unpacked"):
        load(session, "mayorista.xlsx", content.getvalue())


def test_excel_error_cells_are_missing_values_not_prices(session):
    run = load(session, "errores.xls")

    # #REF! and #DIV/0! plus four blank markets; none is read as a price.
    assert (run.prices_loaded, run.rows_rejected, run.values_missing) == (1, 0, 6)
    assert session.query(Price).one().amount == Decimal("20.00")


def test_new_products_take_the_most_common_rubro(session):
    rows = [
        "Leche (Líquida);Avicolas;Mayorista;Fardo/12 Ud;900;Primera;Agosto;2017",
        "Leche (Líquida);Lacteos;Mayorista;Fardo/12 Ud;910;Segunda;Agosto;2017",
        "Leche (Líquida);Lateos;Mayorista;Fardo/12 Ud;920;Tercera;Agosto;2017",
    ]

    load(session, "mayorista.csv", "\r\n".join(rows).encode("latin-1"))

    assert session.query(Product).one().category == "Lacteos"


def wholesale(*rows: str) -> bytes:
    return "\r\n".join(rows).encode("latin-1")


def test_an_isolated_date_typo_takes_the_date_of_its_agreeing_neighbors(session):
    content = wholesale(
        "Papa;Raices-Tuberculos;Mayorista;Quintal;1500;Primera;Junio;2018",
        "Yuca;Raices-Tuberculos;Mayorista;Quintal;1200;3333;Junio;2018",
        "Batata;Raices-Tuberculos;Mayorista;Quintal;1100;Primera;Junio;2018",
        "Ñame;Raices-Tuberculos;Mayorista;Quintal;2000;Segunda;Juio;2018",
        "Ñame;Raices-Tuberculos;Mayorista;Quintal;2100;;;",
        "Ñame;Raices-Tuberculos;Mayorista;Quintal;2200;Segunda;Junio;2018",
    )

    run = load(session, "mayorista.csv", content)

    yuca = session.query(Price).join(Product).filter(Product.name == "Yuca").one()
    assert yuca.observed_on == date(2018, 6, 1)
    assert run.rows_corrected == 1
    # A row without any date keeps its error; a typo between different weeks is not guessed.
    assert reasons(run) == {RejectionReason.INVALID_DATE: 1}
    names = Counter(price.product.name for price in session.query(Price))
    assert names == {"Papa": 1, "Yuca": 1, "Batata": 1, "Ñame": 2}


def test_a_name_listed_twice_in_one_week_is_a_second_product(session):
    eggplant = "Berenjena (Pompadur), primera;Legumbres-Hortalizas;Mayorista;Saco/125 lb"
    content = wholesale(
        f"{eggplant};1300;Primera;Agosto;2017",
        f"{eggplant};900;Primera;Agosto;2017",
        f"{eggplant};1250;Segunda;Agosto;2017",
        f"{eggplant};950;Segunda;Agosto;2017",
    )

    run = load(session, "mayorista.csv", content)

    prices = {
        (price.product.name, price.observed_on.day): price.amount for price in session.query(Price)
    }
    assert prices == {
        ("Berenjena (Pompadur), primera", 1): Decimal("1300.00"),
        ("Berenjena (Pompadur), primera (fila 2)", 1): Decimal("900.00"),
        ("Berenjena (Pompadur), primera", 8): Decimal("1250.00"),
        ("Berenjena (Pompadur), primera (fila 2)", 8): Decimal("950.00"),
    }
    assert run.rows_rejected == 0


def test_a_lasting_change_of_unit_starts_a_new_series_and_stray_units_are_rejected(session):
    weeks = ["Primera", "Segunda", "Tercera", "Cuarta"]
    content = wholesale(
        *[f"Pepino;Legumbres-Hortalizas;Mayorista;Saco/90 lb;600;{w};Enero;2024" for w in weeks],
        *[f"Pepino;Legumbres-Hortalizas;Mayorista;Ciento;900;{w};Febrero;2024" for w in weeks],
        "Pepino;Legumbres-Hortalizas;Mayorista;Litro;50;Primera;Marzo;2024",
    )

    run = load(session, "mayorista.csv", content)

    bases = Counter(price.unit.base_unit for price in session.query(Price))
    assert bases == {BaseUnit.POUND: 4, BaseUnit.UNIT: 4}
    assert reasons(run) == {RejectionReason.UNIT_MISMATCH: 1}


def test_names_where_the_source_wrote_a_semicolon_for_enye_are_repaired(session):
    content = wholesale(
        '"Aguacate (Semíl-34), primera, peque;o";Frutas;Mayorista ;Ciento;4000;Tercera;Abril;2026'
    )

    load(session, "mayorista.csv", content)

    assert session.query(Product).one().name == "Aguacate (Semíl-34), primera, pequeño"


def test_a_repeat_far_inside_a_mislabeled_block_stays_a_duplicate(session):
    # Two weeks under one label: every product appears twice, but not side by side.
    names = ["Papa", "Yuca", "Batata", "Ñame"]
    week = [f"{name};Raices-Tuberculos;Mayorista;Quintal;1000;Tercera;Mayo;2023" for name in names]

    run = load(session, "mayorista.csv", wholesale(*week, *week))

    assert session.query(Product).count() == 4
    assert reasons(run) == {RejectionReason.DUPLICATE_PRICE: 4}


def test_a_valid_date_between_rows_of_another_week_is_kept(session):
    content = wholesale(
        "Papa;Raices-Tuberculos;Mayorista;Quintal;1500;Primera;Agosto;2017",
        "Yuca;Raices-Tuberculos;Mayorista;Quintal;1200;Segunda;Agosto;2017",
        "Batata;Raices-Tuberculos;Mayorista;Quintal;1100;Primera;Agosto;2017",
    )

    run = load(session, "mayorista.csv", content)

    yuca = session.query(Price).join(Product).filter(Product.name == "Yuca").one()
    assert yuca.observed_on == date(2017, 8, 8)
    assert run.rows_corrected == 0


def test_rows_without_a_name_or_too_short_are_rejected_not_repaired(session):
    content = wholesale(
        "Papa;Raices-Tuberculos;Mayorista;Quintal;1500;Primera;Agosto;2017",
        ";Raices-Tuberculos;Mayorista;Quintal;1200;Primera;Agosto;2017",
        "Yuca;Raices-Tuberculos;Mayorista;Quintal;1200",
        ";Raices-Tuberculos;Mayorista;;;Quinta;Noviembre;2023",
        "Batata;Raices-Tuberculos;Mayorista;Quintal;1100;Primera;Agosto;2017",
    )

    run = load(session, "mayorista.csv", content)

    # A row with neither name nor price has nothing to load: a missing value, not an error.
    assert reasons(run) == {RejectionReason.UNEXPECTED_LAYOUT: 2}
    assert run.values_missing == 1
    assert sorted(product.name for product in session.query(Product)) == ["Batata", "Papa"]


def test_a_short_row_with_a_partial_date_is_rejected_not_repaired(session):
    content = wholesale(
        "Papa;Raices-Tuberculos;Mayorista;Quintal;1500;Primera;Agosto;2017",
        "Yuca;Raices-Tuberculos;Mayorista;Quintal;1200;3333",
        "Batata;Raices-Tuberculos;Mayorista;Quintal;1100;Primera;Agosto;2017",
    )

    run = load(session, "mayorista.csv", content)

    assert reasons(run) == {RejectionReason.UNEXPECTED_LAYOUT: 1}
    assert run.rows_corrected == 0


def test_an_empty_name_cell_in_a_workbook_is_rejected(session):
    workbook = openpyxl.Workbook()
    workbook.active.append(
        ["Papa", "Raices", "Mayorista", "Quintal", 1500, "Primera", "Agosto", 2017]
    )
    workbook.active.append(
        [None, "Raices", "Mayorista", "Quintal", 1200, "Primera", "Agosto", 2017]
    )
    content = io.BytesIO()
    workbook.save(content)

    run = load(session, "mayorista.xlsx", content.getvalue())

    assert reasons(run) == {RejectionReason.UNEXPECTED_LAYOUT: 1}
    assert [product.name for product in session.query(Product)] == ["Papa"]
